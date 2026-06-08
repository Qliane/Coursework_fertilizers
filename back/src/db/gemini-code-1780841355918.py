import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = docx.Document()
    
    # Настройка полей страницы
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Вспомогательные функции для стилизации текста
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = RGBColor(51, 102, 153)
        return p

    def add_body_text(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        return p
        
    def add_formula(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Cambria Math'
        run.font.size = Pt(11)
        return p

    def add_code_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        return p

    # Заголовок документа
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("Примеры формул реляционной алгебры и исчисления")
    run.font.name = 'Arial'
    run.font.size = Pt(18)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    # --- ЗАПРОС А ---
    add_heading_1("Запрос А. Список водителей, работающих на заданного партнёра")
    
    add_heading_2("Естественный язык")
    add_body_text("Выбрать всех водителей (их идентификатор, идентификатор партнёра, номер лицензии, категории, имя, фамилию, отчество, идентификатор роли и название роли), которые принадлежат партнёру с заданным идентификатором (P).")
    
    add_heading_2("Реляционная алгебра")
    add_formula("DriverFull ← DRIVER ⋈_{DRIVER.USER_ID = USER.USER_ID} USER")
    add_formula("ResultTemp ← DriverFull ⋈_{DriverFull.ROLE_ID = ROLE.ROLE_ID} ROLE")
    add_formula("Result ← 𝜋_{USER_ID → userId, PARTNER_ID → partnerId, DRIVER_LICENSE → license, DRIVER_CATEGORIES → categories, USER_NAME → name, USER_SECONDNAME → surname, USER_PATRONYMIC → patronymic, ROLE_ID → roleId, ROLE_NAME → roleName} ( 𝜎_{PARTNER_ID = P} (ResultTemp) )")
    
    add_heading_2("Реляционное исчисление (переменные-кортежи)")
    add_body_text("Пусть d – кортеж отношения DRIVER, u – кортеж отношения USER, r – кортеж отношения ROLE.\nЗадан параметр P (идентификатор партнёра).")
    add_body_text("Результат: множество кортежей res таких, что:")
    
    calc_a = (
        "∃ d ∃ u ∃ r (\n"
        "  res.userId = d.USER_ID ∧\n"
        "  res.partnerId = d.PARTNER_ID ∧\n"
        "  res.license = d.DRIVER_LICENSE ∧\n"
        "  res.categories = d.DRIVER_CATEGORIES ∧\n"
        "  res.name = u.USER_NAME ∧\n"
        "  res.surname = u.USER_SECONDNAME ∧\n"
        "  res.patronymic = u.USER_PATRONYMIC ∧\n"
        "  res.roleId = u.ROLE_ID ∧\n"
        "  res.roleName = r.ROLE_NAME ∧\n"
        "  d.PARTNER_ID = P ∧\n"
        "  u.USER_ID = d.USER_ID ∧\n"
        "  r.ROLE_ID = u.ROLE_ID )"
    )
    add_formula(calc_a)
    
    add_heading_2("SQL")
    sql_a = (
        "SELECT\n"
        "    d.USER_ID as userId,\n"
        "    d.PARTNER_ID as partnerId,\n"
        "    d.DRIVER_LICENSE as license,\n"
        "    d.DRIVER_CATEGORIES as categories,\n"
        "    u.USER_NAME as name,\n"
        "    u.USER_SECONDNAME as surname,\n"
        "    u.USER_PATRONYMIC as patronymic,\n"
        "    u.ROLE_ID as roleId,\n"
        "    TRIM(r.ROLE_NAME) as roleName\n"
        "FROM DRIVER d\n"
        "INNER JOIN \"USER\" u ON d.USER_ID = u.USER_ID\n"
        "INNER JOIN ROLE r ON u.ROLE_ID = r.ROLE_ID\n"
        "WHERE d.PARTNER_ID = $1\n"
        "ORDER BY u.USER_SECONDNAME, u.USER_NAME;"
    )
    add_code_block(sql_a)
    
    add_heading_2("Пример выполнения (для P = 10)")
    table_a = doc.add_table(rows=3, cols=9)
    table_a.style = 'Light Shading Accent 1'
    
    headers = ["userId", "partnerId", "license", "categories", "name", "surname", "patronymic", "roleId", "roleName"]
    for i, h in enumerate(headers):
        table_a.rows[0].cells[i].text = h
        
    data1 = ["101", "10", "1234 567890", "C, E", "Иван", "Петров", "Иванович", "3", "DRIVER"]
    for i, d in enumerate(data1):
        table_a.rows[1].cells[i].text = d
        
    data2 = ["102", "10", "9876 543210", "C", "Сергей", "Сидоров", "Алексеевич", "3", "DRIVER"]
    for i, d in enumerate(data2):
        table_a.rows[2].cells[i].text = d

    # --- ЗАПРОС Б ---
    doc.add_page_break()
    add_heading_1("Запрос Б. Список транспортных средств заданного партнёра")
    
    add_heading_2("Естественный язык")
    add_body_text("Выбрать все автомобили (идентификатор, идентификатор партнёра, регистрационный знак, тип, грузоподъёмность), принадлежащие партнёру с заданным идентификатором (P).")
    
    add_heading_2("Реляционная алгебра")
    add_formula("Result ← 𝜋_{VEHICLE_ID → id, PARTNER_ID → partnerId, VEHICLE_REGISTRATION_MARK → registrationMark, VEHICLE_TYPE → type, VEHICLE_CAPACITY → capacity} ( 𝜎_{PARTNER_ID = P} (VEHICLE) )")
    
    add_heading_2("Реляционное исчисление (переменные-кортежи)")
    add_body_text("Пусть v – кортеж отношения VEHICLE.\nЗадан параметр P (идентификатор партнёра).")
    add_body_text("Результат: множество кортежей res таких, что:")
    
    calc_b = (
        "∃ v (\n"
        "  res.id = v.VEHICLE_ID ∧\n"
        "  res.partnerId = v.PARTNER_ID ∧\n"
        "  res.registrationMark = v.VEHICLE_REGISTRATION_MARK ∧\n"
        "  res.type = v.VEHICLE_TYPE ∧\n"
        "  res.capacity = v.VEHICLE_CAPACITY ∧\n"
        "  v.PARTNER_ID = P )"
    )
    add_formula(calc_b)
    
    add_heading_2("SQL")
    sql_b = (
        "SELECT\n"
        "    VEHICLE_ID as id,\n"
        "    PARTNER_ID as partnerId,\n"
        "    TRIM(VEHICLE_REGISTRATION_MARK) as registrationMark,\n"
        "    VEHICLE_TYPE as type,\n"
        "    VEHICLE_CAPACITY as capacity\n"
        "FROM VEHICLE\n"
        "WHERE PARTNER_ID = $1\n"
        "ORDER BY VEHICLE_ID;"
    )
    add_code_block(sql_b)
    
    add_heading_2("Пример выполнения (для P = 10)")
    table_b = doc.add_table(rows=3, cols=5)
    table_b.style = 'Light Shading Accent 1'
    
    headers_b = ["id", "partnerId", "registrationMark", "type", "capacity"]
    for i, h in enumerate(headers_b):
        table_b.rows[0].cells[i].text = h
        
    data1_b = ["51", "10", "А123ВС 77", "Фургон", "15000"]
    for i, d in enumerate(data1_b):
        table_b.rows[1].cells[i].text = d
        
    data2_b = ["52", "10", "В456ОЕ 99", "Рефрижератор", "12000"]
    for i, d in enumerate(data2_b):
        table_b.rows[2].cells[i].text = d

    doc.save("Relational_Queries.docx")
    print("Файл 'Relational_Queries.docx' успешно создан!")

if __name__ == "__main__":
    create_document()